"""Real OOXML through the actual Loader dispatch; no app, network or encoding."""

import ast
import asyncio
import io
import logging
import os
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi import HTTPException
from langchain_core.documents import Document
from open_webui.retrieval.loaders.local import DocxLoader, HTMLLoader, TextLoader
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]


def fixture_file(root, extension, *, pictures=True):
    image = io.BytesIO()
    Image.new('RGB', (40, 30), 'red').save(image, format='PNG')
    path = root / f'office-slim.{extension}'
    if extension == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches

        document = Presentation()
        for title in ('ORBIT-PPTX-731', 'SECOND-SLIDE-42'):
            slide = document.slides.add_slide(document.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = title
            if pictures:
                slide.shapes.add_picture(image, Inches(1), Inches(2))
        # Package numbering is not presentation order.
        ids = document.slides._sldIdLst
        ids.insert(0, ids[-1])
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage

        document = Workbook()
        sheet = document.active
        sheet.title = 'Measurements'
        sheet['A1'] = 'ORBIT-XLSX-914'
        sheet['C1'] = 73
        sheet['D1'] = True
        sheet['E1'] = '=C1*2'
        if pictures:
            sheet.add_image(SheetImage(image), 'C5')
        document.create_sheet('Second sheet')['B3'] = 'SECOND-SHEET-42'
    document.save(path)
    return path


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    monkeypatch.setattr('requests.sessions.Session.request', Mock(side_effect=AssertionError('network forbidden')))


@pytest.fixture
def loader_factory(tmp_path):
    source = ROOT / 'backend/open_webui/retrieval/loaders/main.py'
    nodes = ast.parse(source.read_text()).body
    selected = [
        node for node in nodes if isinstance(node, ast.ClassDef) and node.name in ('Loader', 'CSVLoaderWithSummary')
    ]
    import csv

    scope = dict(
        Document=Document,
        Path=Path,
        os=os,
        zipfile=zipfile,
        asyncio=asyncio,
        csv=csv,
        USE_SLIM=True,
        HTTPException=HTTPException,
        DocxLoader=DocxLoader,
        HTMLLoader=HTMLLoader,
        TextLoader=TextLoader,
        known_archive_ext={'docx', 'pptx', 'xlsx'},
        known_archive_content_types=set(),
        known_source_ext=['py'],
        log=logging.getLogger('office-test'),
        ftfy=SimpleNamespace(fix_text=lambda value, **kwargs: value),
    )
    exec(compile(ast.Module(body=selected, type_ignores=[]), str(source), 'exec'), scope)

    def build(**overrides):
        config = {
            'rag.content_extraction_engine': 'mineru',
            'rag.mineru_file_extensions': ['pdf'],
            'rag.file.max_size': 100,
            'rag.extract_document_image_assets': True,
            'rag.document_image_asset_root': str(tmp_path / 'assets'),
        }
        config.update(overrides)
        return scope['Loader'](config), scope

    return build


@pytest.mark.parametrize('extension', ['pptx', 'xlsx'])
def test_real_slim_dispatch_body_images_and_anchors(tmp_path, loader_factory, extension):
    path = fixture_file(tmp_path, extension)
    loader, _ = loader_factory()
    docs = loader.load(path.name, 'application/octet-stream', str(path))
    assert len(docs) == 2
    assert all('page' not in doc.metadata and 'page_index' not in doc.metadata for doc in docs)
    assets = docs[0].metadata['document_image_assets']
    assert not docs[0].metadata['document_image_assets_skipped']
    assert all(Path(asset['storage_uri']).is_file() for asset in assets)
    if extension == 'pptx':
        assert 'SECOND-SLIDE-42' in docs[0].page_content
        assert 'ORBIT-PPTX-731' in docs[1].page_content
        assert [doc.metadata['anchor']['slide_index'] for doc in docs] == [0, 1]
        assert docs[0].metadata['anchor']['part'] == 'ppt/slides/slide2.xml'
        assert {asset['anchor']['slide_index'] for asset in assets} == {0, 1}
    else:
        assert 'A1: ORBIT-XLSX-914' in docs[0].page_content
        assert 'C1: 73' in docs[0].page_content and 'D1: TRUE' in docs[0].page_content
        assert 'E1: =C1*2' in docs[0].page_content
        assert 'B3: SECOND-SHEET-42' in docs[1].page_content
        assert [doc.metadata['anchor']['sheet_name'] for doc in docs] == ['Measurements', 'Second sheet']
        assert assets[0]['anchor']['sheet_name'] == 'Measurements'
        assert assets[0]['anchor']['row_index'] == 4 and assets[0]['anchor']['col_index'] == 2


def test_unchanged_docx_csv_text_and_pdf_dispatch(tmp_path, loader_factory, monkeypatch):
    from docx import Document as WordDocument

    path = tmp_path / 'ordinary.docx'
    doc = WordDocument()
    doc.add_paragraph('Existing DOCX text')
    doc.save(path)
    loader, scope = loader_factory()
    assert 'Existing DOCX text' in loader.load(path.name, '', str(path))[0].page_content
    path = tmp_path / 'ordinary.csv'
    path.write_text('name,value\nalpha,17\n')
    monkeypatch.setattr(loader, '_detect_text_encoding', lambda _: 'utf-8')
    assert 'value: 17' in loader.load(path.name, '', str(path))[0].page_content
    path = tmp_path / 'ordinary.txt'
    path.write_text('Unchanged text')
    assert loader.load(path.name, 'text/plain', str(path))[0].page_content == 'Unchanged text'
    scope['MinerULoader'] = Mock()
    loader.config.update(
        {
            'rag.mineru_api_timeout': 300,
            'rag.mineru_api_mode': 'cloud',
            'rag.mineru_api_url': 'https://example.invalid',
            'rag.mineru_api_key': '',
            'rag.mineru_params': {},
        }
    )
    scope['MINERU_MAX_MARKDOWN_BYTES'] = 1000
    loader._get_loader('ordinary.pdf', 'application/pdf', str(path))
    assert scope['MinerULoader'].call_count == 1


@pytest.mark.parametrize('extension', ['ppt', 'xls'])
def test_binary_office_remains_external_only(tmp_path, loader_factory, extension):
    path = tmp_path / ('binary.' + extension)
    path.write_bytes(b'legacy binary not OOXML')
    loader, _ = loader_factory()
    with pytest.raises(HTTPException) as error:
        loader.load(path.name, '', str(path))
    assert error.value.status_code == 503


def test_shared_strings_formula_cache_and_no_phonetic_duplication(tmp_path):
    from open_webui.retrieval.loaders.office_text import OfficeTextLoader

    path = tmp_path / 'shared.xlsx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr(
            'xl/workbook.xml', '<workbook xmlns:r="urn:rel"><sheets><sheet name="Data" r:id="r1"/></sheets></workbook>'
        )
        archive.writestr(
            'xl/_rels/workbook.xml.rels',
            '<Relationships><Relationship Id="r1" Target="worksheets/data.xml"/></Relationships>',
        )
        archive.writestr(
            'xl/sharedStrings.xml',
            '<sst><si><r><t>Rich </t></r><r><t>text</t></r><rPh><t>phonetic</t></rPh></si></sst>',
        )
        archive.writestr(
            'xl/worksheets/data.xml',
            '<worksheet><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c>'
            '<c r="D1"><f>2*3</f><v>6</v></c></row></sheetData></worksheet>',
        )
    docs = OfficeTextLoader(path, 'xlsx').load()
    assert docs[0].page_content == 'A1: Rich text\tD1: 6'
    assert docs[0].metadata['anchor'] == {'sheet_name': 'Data', 'part': 'xl/worksheets/data.xml'}


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-16'])
def test_xml_entities_rejected_in_text_parts(tmp_path, encoding):
    from open_webui.retrieval.loaders.office_text import OfficeTextLoader

    path = tmp_path / 'entity.xlsx'
    with zipfile.ZipFile(path, 'w') as archive:
        text = (
            '<?xml version="1.0" encoding="' + encoding + '"?>'
            '<!DOCTYPE workbook [<!ENTITY x "expansion">]><workbook>&x;</workbook>'
        )
        archive.writestr('xl/workbook.xml', text.encode(encoding))
    with pytest.raises(ValueError, match='entities_forbidden'):
        OfficeTextLoader(path, 'xlsx').load()


def test_metadata_only_image_package_preserves_existing_hook(tmp_path, loader_factory):
    path = fixture_file(tmp_path, 'xlsx')
    with zipfile.ZipFile(path) as original:
        data = {name: original.read(name) for name in original.namelist()}
    for name in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'):
        data[name] = b'<worksheet/>'
    with zipfile.ZipFile(path, 'w') as archive:
        for name, value in data.items():
            archive.writestr(name, value)
    loader, _ = loader_factory()
    docs = loader.load(path.name, '', str(path))
    assert len(docs) == 1 and docs[0].metadata['_metadata_only']
    assert len(docs[0].metadata['document_image_assets']) == 1


def test_image_extraction_toggle_still_respected(tmp_path, loader_factory):
    path = fixture_file(tmp_path, 'pptx')
    loader, _ = loader_factory(**{'rag.extract_document_image_assets': False})
    docs = asyncio.run(loader.aload(path.name, '', str(path)))
    assert len(docs) == 2 and all('document_image_assets' not in doc.metadata for doc in docs)


@pytest.mark.parametrize('kind', ['slide_index', 'sheet_name'])
def test_markdown_min_size_merge_must_not_cross_office_locations(kind):
    source = ROOT / 'backend/open_webui/routers/retrieval.py'
    function = next(
        node
        for node in ast.parse(source.read_text()).body
        if isinstance(node, ast.FunctionDef) and node.name == 'can_merge_chunks'
    )
    scope = {'Document': Document}
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(source), 'exec'), scope)
    first = Document(page_content='one', metadata={'source': 'same', 'file_id': 'same', 'anchor': {kind: 0}})
    second = Document(page_content='two', metadata={'source': 'same', 'file_id': 'same', 'anchor': {kind: 1}})
    assert not scope['can_merge_chunks'](first, second)
    second.metadata['anchor'] = first.metadata['anchor']
    assert scope['can_merge_chunks'](first, second)
    first.metadata.pop('anchor')
    second.metadata.pop('anchor')
    assert scope['can_merge_chunks'](first, second)
