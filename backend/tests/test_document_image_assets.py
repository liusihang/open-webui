"""Offline contracts: actual OOXML/PDF files and parser-response fixtures.

No application import, DB, GPU, OCR, embedding, or real HTTP request.
"""

import ast
import base64
import io
import json
import zipfile
from pathlib import Path
from unittest.mock import Mock

import pytest
from langchain_core.documents import Document
from open_webui.retrieval.document_image_assets import (
    ImageAssetStore,
    download_service_resource,
    extract_markdown_image_assets,
)
from open_webui.retrieval.loaders.mineru import MinerULoader
from open_webui.retrieval.loaders.office_image_assets import extract_office_image_assets
from open_webui.retrieval.loaders.paddleocr_vl import PaddleOCRVLLoader, _page_index
from open_webui.retrieval.loaders.pdf_image_assets import extract_pdf_image_assets
from PIL import Image

FIXTURE = Path(__file__).parent / 'fixtures/document_image_assets/paddleocr_layout.json'


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def blocked(*args, **kwargs):
        raise AssertionError('Unexpected network request')

    monkeypatch.setattr('requests.sessions.Session.request', blocked)


@pytest.fixture
def png():
    return base64.b64decode(
        json.loads(FIXTURE.read_text())['result']['layoutParsingResults'][1]['markdown']['images']['imgs/figure.png']
    )


@pytest.fixture
def source(tmp_path):
    path = tmp_path / 'source.pdf'
    path.write_bytes(b'%PDF-fixture-not-parsed-by-mocked-OCR')
    return path


def response(payload, status=200):
    result = Mock(status_code=status)
    result.json.return_value = payload
    result.__enter__ = Mock(return_value=result)
    result.__exit__ = Mock(return_value=False)
    return result


def test_paddle_layout_actual_image_bytes_and_page(monkeypatch, tmp_path, source, png):
    post = Mock(return_value=response(json.loads(FIXTURE.read_text())))
    monkeypatch.setattr('requests.post', post)
    loader = PaddleOCRVLLoader(
        'https://ocr.example', 'test-token', str(source), extract_images=True, asset_root=tmp_path / 'assets'
    )
    docs = loader.load()
    assert len(docs) == 2
    image = docs[1].metadata['document_image_assets'][0]
    assert (image['page_index'], image['width'], image['height']) == (1, 24, 16)
    assert image['caption'] == 'Existing red figure caption'
    assert Path(image['storage_uri']).read_bytes() == png
    assert image['mime_type'] == 'image/png'
    assert len(list((tmp_path / 'assets').rglob('*.png'))) == 1  # not outputImages
    assert post.call_count == 1  # no secondary OCR
    assert post.call_args.kwargs['json']['returnMarkdownImages'] is True


def test_paddle_jobs_jsonl_contract(monkeypatch, tmp_path, source):
    post = Mock(return_value=response({'code': 0, 'data': {'jobId': 'test-job'}}))
    get = Mock(
        return_value=response(
            {'data': {'state': 'done', 'resultUrl': {'jsonUrl': 'https://assets.example/results.jsonl'}}}
        )
    )
    rows = json.loads(FIXTURE.read_text())
    download = Mock(return_value=(json.dumps(rows) + '\n').encode())
    monkeypatch.setattr('requests.post', post)
    monkeypatch.setattr('requests.get', get)
    monkeypatch.setattr('open_webui.retrieval.loaders.paddleocr_vl.download_service_resource', download)
    loader = PaddleOCRVLLoader(
        'https://ocr.example/api/v2/ocr',
        'test-token',
        str(source),
        api_mode='jobs',
        extract_images=True,
        asset_root=tmp_path / 'assets',
        allowed_remote_origins=['https://assets.example'],
    )
    docs = loader.load()
    assert post.call_args.args[0] == 'https://ocr.example/api/v2/ocr/jobs'
    assert 'file' in post.call_args.kwargs['files']
    assert json.loads(post.call_args.kwargs['data']['optionalPayload'])['restructurePages'] is False
    assert docs[1].metadata['document_image_assets'][0]['page_index'] == 1
    assert download.call_args.kwargs['allowed_origins'] == (loader.api_url, 'https://assets.example')


def test_paddle_failure_not_invented_text(monkeypatch, source):
    monkeypatch.setattr('requests.post', Mock(return_value=response({'private-error': 'secret'}, 503)))
    with pytest.raises(RuntimeError, match='HTTP 503'):
        PaddleOCRVLLoader('https://ocr.example', 'test-token', str(source)).load()


@pytest.mark.parametrize(
    'row,expected', [({'pageIndex': 0}, 0), ({'pageNo': 2}, 1), ({'prunedResult': {'page_index': 4}}, 4), ({}, 7)]
)
def test_real_page_conventions(row, expected):
    assert _page_index(row, 7) == expected


def test_store_once_repeated_occurrence_keeps_anchors(tmp_path, png):
    store = ImageAssetStore(tmp_path / 'assets', '../../source-id')
    a = store.put(png, backend='test', origin_reference='a', page_index=0)
    b = store.put(png, backend='test', origin_reference='b', page_index=1)
    assert a['storage_uri'] == b['storage_uri']
    assert b['page_index'] == 1
    assert len(list((tmp_path / 'assets').rglob('*.png'))) == 1
    assert Path(a['storage_uri']).is_relative_to(tmp_path / 'assets')


def test_untrusted_reference_bad_image_and_missing_resource(tmp_path, png):
    store = ImageAssetStore(tmp_path, 'source')
    extraction = extract_markdown_image_assets(
        {
            'text': 'real text',
            'images': {
                '../escape.png': base64.b64encode(png).decode(),
                'html.png': base64.b64encode(b'<html>not an image</html>').decode(),
                'external.png': 'http://169.254.169.254/latest/meta-data',
            },
        },
        store=store,
        page_index=0,
        backend='test',
        resource_base_url='https://ocr.example',
        allowed_origins=('https://ocr.example',),
    )
    assert extraction.assets == []
    assert len(extraction.skipped) == 3
    missing = extract_markdown_image_assets(
        {'text': '![x](https://evil.example/a.png)'},
        store=store,
        page_index=0,
        backend='test',
        resource_base_url='https://ocr.example',
        allowed_origins=(),
    )
    assert missing.skipped[0]['reason'] == 'parser_image_resources_missing'
    assert not list(tmp_path.rglob('*.png'))


def test_presigned_resource_download_is_materialized(monkeypatch, tmp_path, png):
    fetch = Mock(return_value=png)
    monkeypatch.setattr('open_webui.retrieval.document_image_assets.download_service_resource', fetch)
    extraction = extract_markdown_image_assets(
        {'images': {'images/a.png': 'https://assets.example/a?signature=private'}},
        store=ImageAssetStore(tmp_path, 'file'),
        page_index=0,
        backend='paddleocr-vl',
        resource_base_url='https://ocr.example',
        allowed_origins=('https://assets.example',),
    )
    assert Path(extraction.assets[0]['storage_uri']).read_bytes() == png
    assert 'signature' not in json.dumps(extraction.assets)


def test_raw_jpeg_base64_is_not_mistaken_for_relative_url(tmp_path):
    image = io.BytesIO()
    Image.new('RGB', (12, 8), 'blue').save(image, format='JPEG')
    encoded = base64.b64encode(image.getvalue()).decode()
    assert encoded.startswith('/9j/')
    extraction = extract_markdown_image_assets(
        {'images': {'image.jpg': encoded}},
        store=ImageAssetStore(tmp_path, 'file'),
        page_index=0,
        backend='test',
        resource_base_url='https://ocr.example',
        allowed_origins=(),
    )
    assert len(extraction.assets) == 1
    assert extraction.assets[0]['mime_type'] == 'image/jpeg'


def test_failed_image_download_attempt_budget_is_per_document(monkeypatch, tmp_path):
    fetch = Mock(side_effect=ValueError('resource_unavailable'))
    monkeypatch.setattr('open_webui.retrieval.document_image_assets.download_service_resource', fetch)
    store = ImageAssetStore(tmp_path, 'file', max_images=1)
    for page_index in (0, 1):
        extract_markdown_image_assets(
            {'images': {'a.png': 'https://ocr.example/a'}},
            store=store,
            page_index=page_index,
            backend='test',
            resource_base_url='https://ocr.example',
            allowed_origins=('https://ocr.example',),
        )
    assert fetch.call_count == 1


def test_remote_origin_redirect_and_size_limits(monkeypatch):
    with pytest.raises(ValueError, match='origin_not_allowed'):
        download_service_resource('https://evil.example/a', allowed_origins=('https://ocr.example',), max_bytes=4)
    session = Mock()
    session.__enter__ = Mock(return_value=session)
    session.__exit__ = Mock(return_value=False)
    reply = response({}, 302)
    session.get.return_value = reply
    monkeypatch.setattr('requests.Session', Mock(return_value=session))
    with pytest.raises(ValueError, match='redirect_not_allowed'):
        download_service_resource('https://ocr.example/a', allowed_origins=('https://ocr.example',), max_bytes=4)
    reply.status_code, reply.headers = 200, {}
    reply.iter_content.return_value = [b'123', b'456']
    with pytest.raises(ValueError, match='too_large'):
        download_service_resource('https://ocr.example/a', allowed_origins=('https://ocr.example',), max_bytes=4)
    assert session.get.call_args.kwargs['allow_redirects'] is False
    assert session.trust_env is False
    assert 'headers' not in session.get.call_args.kwargs  # no token leakage


def test_docx_picture_has_paragraph_not_fake_page(tmp_path, png):
    from docx import Document as WordDocument

    path = tmp_path / 'embedded.docx'
    doc = WordDocument()
    doc.add_paragraph('The figure introduction')
    paragraph = doc.add_paragraph('Figure paragraph')
    paragraph.add_run().add_picture(io.BytesIO(png))
    doc.save(path)
    result = extract_office_image_assets(path, asset_root=tmp_path / 'assets')
    assert not result.skipped
    assert len(result.assets) == 1
    asset = result.assets[0]
    assert asset['page_index'] is None
    assert asset['anchor']['part'] == 'word/document.xml'
    assert asset['anchor']['paragraph_index'] == 1
    assert 'Figure paragraph' in asset['surrounding_text']


def test_pptx_same_picture_on_two_slides_stored_once(tmp_path, png):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / 'embedded.pptx'
    presentation = Presentation()
    for _ in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1))
    presentation.save(path)
    result = extract_office_image_assets(path, asset_root=tmp_path / 'assets')
    assert len(result.assets) == 2
    assert {asset['anchor']['slide_index'] for asset in result.assets} == {0, 1}
    assert all(asset['page_index'] is None for asset in result.assets)
    assert len({asset['storage_uri'] for asset in result.assets}) == 1


def test_xlsx_picture_real_sheet_cell(tmp_path, png):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorksheetImage

    path = tmp_path / 'embedded.xlsx'
    workbook = Workbook()
    workbook.active.title = 'Measurements'
    workbook.active.add_image(WorksheetImage(io.BytesIO(png)), 'C5')
    workbook.save(path)
    result = extract_office_image_assets(path, asset_root=tmp_path / 'assets')
    assert len(result.assets) == 1
    anchor = result.assets[0]['anchor']
    assert (anchor['sheet_name'], anchor['row_index'], anchor['col_index']) == ('Measurements', 4, 2)
    assert result.assets[0]['page_index'] is None


def test_pdf_actual_embedded_image_no_ocr(tmp_path, png):
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    path = tmp_path / 'embedded.pdf'
    canvas = Canvas(str(path))
    canvas.drawString(30, 40, 'First text page')
    canvas.showPage()
    canvas.drawImage(ImageReader(io.BytesIO(png)), 30, 40, 120, 80)
    canvas.save()
    result = extract_pdf_image_assets(
        path, asset_root=tmp_path / 'assets', page_text_by_index={1: 'Existing nearby text'}
    )
    assert not result.skipped
    assert len(result.assets) == 1
    assert result.assets[0]['page_index'] == 1
    assert result.assets[0]['surrounding_text'] == 'Existing nearby text'


def test_empty_and_damaged_documents_and_path_traversal(tmp_path, png):
    from docx import Document as WordDocument

    path = tmp_path / 'empty.docx'
    WordDocument().save(path)
    assert extract_office_image_assets(path, asset_root=tmp_path / 'assets').assets == []
    path.write_bytes(b'broken zip')
    assert extract_office_image_assets(path, asset_root=tmp_path / 'assets').skipped
    path.write_bytes(b'broken pdf')
    assert extract_pdf_image_assets(path, asset_root=tmp_path / 'assets').skipped
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('../escape.png', png)
    result = extract_office_image_assets(path, asset_root=tmp_path / 'assets')
    assert result.assets == []
    assert result.skipped[0]['reason'] == 'unsafe_image_reference'


def test_asset_budgets_mime_sniffing(tmp_path, png):
    store = ImageAssetStore(tmp_path, 'file', max_images=1)
    result = store.put(png, backend='test', origin_reference='fake.jpg')
    assert result['mime_type'] == 'image/png'
    with pytest.raises(ValueError, match='budget'):
        store.put(png, backend='test', origin_reference='other.png')
    with pytest.raises(ValueError, match='size'):
        ImageAssetStore(tmp_path, 'large', max_image_bytes=1).put(png, backend='test', origin_reference='x')


def test_asset_source_directory_cannot_escape_via_symlink(tmp_path, png):
    store = ImageAssetStore(tmp_path / 'assets', 'file')
    store.root.parent.mkdir()
    outside = tmp_path / 'outside'
    outside.mkdir()
    store.root.symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match='symlink'):
        store.put(png, backend='test', origin_reference='x')
    assert not list(outside.iterdir())


def test_mineru_inline_images_and_existing_page_caption(tmp_path, source, png):
    loader = MinerULoader(str(source), extract_images=True, asset_root=tmp_path / 'assets')
    loader._materialize_images(
        {'images/a.png': 'data:image/png;base64,' + base64.b64encode(png).decode()},
        [{'img_path': 'images/a.png', 'page_idx': 3, 'image_caption': ['Figure caption']}],
        'Existing markdown',
    )
    asset = loader.image_extraction.assets[0]
    assert (asset['page_index'], asset['caption']) == (3, 'Figure caption')
    assert Path(asset['storage_uri']).read_bytes() == png


def test_mineru_archive_images_persist_after_zip_closes(tmp_path, source, png):
    path = tmp_path / 'result.zip'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('doc/images/a.png', png)
        archive.writestr('doc/doc_content_list.json', json.dumps([{'img_path': 'images/a.png', 'page_idx': 2}]))
    loader = MinerULoader(str(source), extract_images=True, asset_root=tmp_path / 'assets')
    with zipfile.ZipFile(path) as archive:
        loader._materialize_zip_images(archive, 'Existing parsed text')
    assert Path(loader.image_extraction.assets[0]['storage_uri']).read_bytes() == png
    assert loader.image_extraction.assets[0]['page_index'] == 2


def test_mineru_local_load_keeps_text_when_image_is_broken(monkeypatch, tmp_path, source):
    post = Mock(
        return_value=response(
            {
                'results': {
                    'source': {
                        'md_content': 'Existing extracted text',
                        'images': {'images/bad.png': 'broken'},
                        'content_list': '[]',
                    }
                }
            }
        )
    )
    monkeypatch.setattr('requests.post', post)
    loader = MinerULoader(str(source), extract_images=True, asset_root=tmp_path / 'assets')
    docs = loader.load()
    assert docs[0].page_content == 'Existing extracted text'
    assert docs[0].metadata['document_image_assets'] == []
    assert docs[0].metadata['document_image_assets_skipped']
    assert post.call_count == 1
    assert post.call_args.kwargs['data']['return_images'] == 'true'


def test_mineru_cloud_archive_download_extracts_assets_before_temp_cleanup(monkeypatch, tmp_path, source, png):
    content = io.BytesIO()
    with zipfile.ZipFile(content, 'w') as archive:
        archive.writestr('doc/full.md', 'Existing text\n![figure](images/a.png)')
        archive.writestr('doc/images/a.png', png)
        archive.writestr('doc/doc_content_list.json', json.dumps([{'img_path': 'images/a.png', 'page_idx': 1}]))
    reply = response({})
    reply.content = content.getvalue()
    monkeypatch.setattr('requests.get', Mock(return_value=reply))
    loader = MinerULoader(str(source), extract_images=True, asset_root=tmp_path / 'assets')
    text = loader._download_and_extract_zip('https://parser-resource.example/result.zip', 'source.pdf')
    assert 'Existing text' in text
    asset = loader.image_extraction.assets[0]
    assert Path(asset['storage_uri']).read_bytes() == png
    assert asset['page_index'] == 1


def test_paddle_image_only_page_and_partial_image_failure(tmp_path, source):
    loader = PaddleOCRVLLoader(
        'https://ocr.example', 'test-token', str(source), extract_images=True, asset_root=tmp_path / 'assets'
    )
    page = json.loads(FIXTURE.read_text())['result']['layoutParsingResults'][1]
    page['markdown']['text'] = ''
    page['markdown']['images']['bad.png'] = 'not-image-base64'
    docs = loader._build_documents([page], resource_base_url='https://ocr.example')
    assert docs[0].page_content == ''
    assert docs[0].metadata['_metadata_only'] is True
    assert len(docs[0].metadata['document_image_assets']) == 1
    assert len(docs[0].metadata['document_image_assets_skipped']) == 1


def test_docx_xml_entity_and_external_image_not_loaded(tmp_path, png):
    path = tmp_path / 'unsafe.docx'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', '<!DOCTYPE test [<!ENTITY x "bad">]><test>&x;</test>')
        archive.writestr(
            'word/_rels/document.xml.rels',
            '<Relationships><Relationship Id="rId1" Target="media/a.png" /></Relationships>',
        )
        archive.writestr('word/media/a.png', png)
    result = extract_office_image_assets(path, asset_root=tmp_path / 'assets')
    assert result.assets == []
    assert result.skipped[0]['reason'] == 'office_part_unreadable'


def test_loader_attaches_office_and_explicit_local_pdf_only(monkeypatch, tmp_path, source):
    # Compile the actual Loader class without importing app/DB/external SDKs.
    module_path = Path(__file__).parents[1] / 'open_webui/retrieval/loaders/main.py'
    module = ast.parse(module_path.read_text())
    cls = next(node for node in module.body if isinstance(node, ast.ClassDef) and node.name == 'Loader')
    scope = {'Document': Document, 'Path': Path, 'ftfy': Mock(fix_text=lambda value, **kwargs: value)}
    exec(compile(ast.Module(body=[cls], type_ignores=[]), str(module_path), 'exec'), scope)
    loader = scope['Loader'](
        {
            'rag.content_extraction_engine': 'paddleocr_vl',
            'rag.extract_document_image_assets': True,
            'rag.document_image_asset_root': str(tmp_path / 'assets'),
        }
    )
    parsed = [Document(page_content='Text remains usable', metadata={})]
    loader._get_loader = Mock(return_value=Mock(load=Mock(return_value=parsed)))
    local = Mock(side_effect=AssertionError('Must not double extract PDF'))
    monkeypatch.setattr('open_webui.retrieval.loaders.pdf_image_assets.extract_pdf_image_assets', local)
    result = loader.load('test.pdf', 'application/pdf', str(source))
    assert result[0].page_content == 'Text remains usable'
    assert result[0].metadata['document_image_assets_skipped'][0]['reason'] == 'selected_parser_has_no_image_assets'
    assert local.call_count == 0
